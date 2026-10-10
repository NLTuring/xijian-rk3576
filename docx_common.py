# -*- coding: utf-8 -*-
'''Word 文档构建工具模块。

从 build_docx.py 抽出可复用的样式与构建块，供：
  - build_docx.py           项目简介
  - build_bp_outline.py     商业计划书设计纲要
  - build_bp_sample.py      商业计划书样板
共用。统一使用宋体（SimSun）。
'''
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "宋体"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
ACCENT2 = RGBColor(0x2E, 0x74, 0xB5)
DARK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "1F4E79"
ZEBRA_FILL = "EDF3F9"
CALLOUT_FILL = "EEF4FA"

PAGE_W_CM = 21.0
PAGE_H_CM = 29.7
MARGIN_LR_CM = 2.3
TABLE_TOTAL_CM = 16.1


def fit_widths(widths):
    total = float(sum(widths))
    if total <= 0:
        return list(widths)
    k = TABLE_TOTAL_CM / total
    return [w * k for w in widths]


def set_run(run, size=10.5, bold=False, color=None, italic=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color if color is not None else DARK
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font)
    return run


def shade(element, fill):
    pr = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def cell_shade(cell, fill):
    shade(cell._tc, fill)


def cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement("w:" + tag)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tblPr.append(mar)


def no_split(row):
    '''禁止表格行跨页断开，避免出现只剩一行的孤立片段。'''
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    node.set(qn("w:val"), "true")
    trPr.append(node)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trPr.append(node)


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:keepNext"))


def set_table_widths(t, widths_cm):
    '''显式写入 tblGrid / tblW / tblLayout，避免渲染器把表格拉成等宽而溢出。'''
    tbl = t._tbl
    tblPr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        el = tblPr.find(qn(tag))
        if el is not None:
            tblPr.remove(el)

    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(round(sum(widths_cm) * 567))))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(round(w * 567))))
        grid.append(gc)
    tbl.insert(list(tbl).index(tblPr) + 1, grid)

    for row in t.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def init_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.font.color.rgb = DARK
    for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        st.element.rPr.rFonts.set(qn(a), FONT)
    st.paragraph_format.line_spacing = 1.4
    st.paragraph_format.space_after = Pt(4)

    for name, size, color, before, after in (
        ("Heading 1", 15, ACCENT, 16, 8),
        ("Heading 2", 12.5, ACCENT2, 12, 5),
        ("Heading 3", 11, ACCENT2, 9, 4),
    ):
        s = doc.styles[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = color
        for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            s.element.rPr.rFonts.set(qn(a), FONT)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.line_spacing = 1.25
        s.paragraph_format.keep_with_next = True


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(PAGE_W_CM)
    sec.page_height = Cm(PAGE_H_CM)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(MARGIN_LR_CM)
    sec.right_margin = Cm(MARGIN_LR_CM)
    return sec


def add_page_number_footer(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("第 ")
        set_run(r, 9, color=GREY)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        inner = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        node = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            node.set(qn(a), FONT)
        rPr.append(node)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "18")
        rPr.append(sz)
        inner.append(rPr)
        t = OxmlElement("w:t")
        t.text = "1"
        inner.append(t)
        fld.append(inner)
        p._p.append(fld)
        r2 = p.add_run(" 页")
        set_run(r2, 9, color=GREY)


# ---------------------------------------------------------------- 构建块
def H(doc, text, level):
    p = doc.add_paragraph(style="Heading %d" % level)
    set_run(p.add_run(text), {1: 15, 2: 12.5, 3: 11}[level], True,
            {1: ACCENT, 2: ACCENT2, 3: ACCENT2}[level])
    return p


def add_rich(paragraph, text, size=10.5, bold=False, color=None):
    '''把 **粗体** 标记解析为真正的加粗 run，其余按普通文本追加。'''
    import re
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text, re.S):
        if m.start() > pos:
            set_run(paragraph.add_run(text[pos:m.start()]), size, bold=bold, color=color)
        set_run(paragraph.add_run(m.group(1)), size, bold=True,
                color=color if color is not None else ACCENT)
        pos = m.end()
    if pos < len(text):
        set_run(paragraph.add_run(text[pos:]), size, bold=bold, color=color)
    return paragraph


def para(doc, text, indent=True, size=10.5, align=None, after=4, color=None):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if "**" in text:
        add_rich(p, text, size, color=color)
    else:
        set_run(p.add_run(text), size, color=color)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75 + 0.6 * level)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.35
    mark = "● " if level == 0 else "○ "
    if "：" in text and text.index("：") < 26:
        head, tail = text.split("：", 1)
        set_run(p.add_run(mark), 9, color=ACCENT2)
        set_run(p.add_run(head + "："), 10.5, bold=True)
        if "**" in tail:
            add_rich(p, tail)
        else:
            set_run(p.add_run(tail), 10.5)
    else:
        set_run(p.add_run(mark), 9, color=ACCENT2)
        if "**" in text:
            add_rich(p, text)
        else:
            set_run(p.add_run(text), 10.5)
    return p


def table(doc, headers, rows, widths=None, size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t)

    hdr = t.rows[0]
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        set_run(p.add_run(h), size, True, RGBColor(0xFF, 0xFF, 0xFF))
        cell_shade(c, HEADER_FILL)
    repeat_header(hdr)
    no_split(hdr)

    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        no_split(t.rows[-1])
        for i, val in enumerate(row):
            c = cells[i]
            c.text = ""
            for li, line in enumerate(str(val).split("\n")):
                p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                set_run(p.add_run(line), size)
            if ri % 2 == 1:
                cell_shade(c, ZEBRA_FILL)

    if widths:
        set_table_widths(t, fit_widths(widths))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(doc, title_text, lines):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.text = ""
    cell_margins(t, top=120, bottom=120, left=160, right=160)
    cell_shade(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(title_text), 11, True, ACCENT)
    for line in lines:
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        q.paragraph_format.line_spacing = 1.35
        q.paragraph_format.first_line_indent = Pt(21)
        set_run(q.add_run(line), 10.5)
    set_table_widths(t, [TABLE_TOTAL_CM])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(text), 9, color=GREY)
    return p


def add_figure(doc, path, cap, width_cm=15.4):
    import os
    if not os.path.exists(path):
        return None
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after = Pt(8)
    set_run(q.add_run(cap), 9, color=GREY)
    return p


def page_break(doc):
    doc.add_page_break()


def new_document():
    doc = Document()
    init_styles(doc)
    setup_page(doc)
    return doc


def cover(doc, kicker, title_lines, subtitle, info_pairs, title_size=22):
    for _ in range(2):
        p = doc.add_paragraph()
        set_run(p.add_run(""), 12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    set_run(p.add_run(kicker), 12, color=GREY)
    for i, line in enumerate(title_lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10 if i == 0 else 0)
        p.paragraph_format.space_after = Pt(4)
        set_run(p.add_run(line), title_size, True, ACCENT)
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        set_run(p.add_run(subtitle), 10.5, False, ACCENT2)

    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(t, top=70, bottom=70, left=120, right=120)
    for k, v in info_pairs:
        cells = t.add_row().cells
        for idx, txt in enumerate((k, v)):
            c = cells[idx]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.3
            if idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_run(p.add_run(txt), 10, True, ACCENT)
                cell_shade(c, ZEBRA_FILL)
            else:
                set_run(p.add_run(txt), 10)
    set_table_widths(t, [2.6, 6.0])
    page_break(doc)
